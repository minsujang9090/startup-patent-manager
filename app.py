import streamlit as st
import pandas as pd
from datetime import datetime
import time
import altair as alt
import openpyxl

# --- 설정: 변경할 파일명 한 곳에서 관리 ---
PATENT_FILE = "patent_rawfile.xlsx"

# --- 1. 기본 설정 ---
st.set_page_config(page_title="사내 특허 관리 시스템 v3.0", layout="wide")

# --- 2. 유틸리티 함수 ---
def classify_cost_stage(text):
    """비용 구분을 보고 단계를 분류하는 함수"""
    if pd.isna(text): return "기타"
    text = str(text).replace(" ", "")
    
    if any(x in text for x in ['연차', '갱신', '유지']):
        return '유지관리(연차료)'
    elif any(x in text for x in ['등록', '설정', '성공보수']):
        return '등록 단계'
    elif any(x in text for x in ['출원', '심사', '거절', '의견', '보정', 'OA', '포기']):
        return '출원 단계'
    else:
        return '기타'

def format_date_clean(val):
    """날짜에서 시간 정보를 떼고 YYYY-MM-DD 문자열만 반환"""
    if pd.isna(val) or str(val).strip() == '':
        return "-"
    try:
        dt = pd.to_datetime(val)
        return dt.strftime('%Y-%m-%d')
    except:
        return str(val)[:10]

# --- 3. 데이터 로드 함수 ---
@st.cache_data
def load_data():
    try:
        file_path = PATENT_FILE
        df_master = pd.read_excel(file_path, sheet_name='기본정보(Master)', engine='openpyxl')
        
        # [수정] PyArrow 호환성 이슈 해결을 위해 ID/번호 컬럼 강제 문자열 변환
        str_cols = ['가산관리번호', '출원번호', '등록번호']
        for col in str_cols:
            if col in df_master.columns:
                # NaN은 빈 문자열로, 나머지는 문자열로 변환하여 mixed types 방지
                df_master[col] = df_master[col].fillna("").astype(str)

        # [추가] 날짜 컬럼(일자, Date 포함) 강제 datetime 변환
        for col in df_master.columns:
            if '일자' in col or 'Date' in col or 'date' in col.lower():
                df_master[col] = pd.to_datetime(df_master[col], errors='coerce')

        df_schedule = pd.read_excel(file_path, sheet_name='비용관리(Schedule)', engine='openpyxl')
        
        # 날짜 컬럼 형변환 (에러 발생 시 NaT로 처리)
        df_schedule['기한(Due Date)'] = pd.to_datetime(df_schedule['기한(Due Date)'], errors='coerce')
        # [중요] 실제납부일도 날짜형으로 변환 시도
        df_schedule['실제납부일'] = pd.to_datetime(df_schedule['실제납부일'], errors='coerce')
        # [추가] 회신요청기한 날짜 변환
        if '회신요청기한' in df_schedule.columns:
            df_schedule['회신요청기한'] = pd.to_datetime(df_schedule['회신요청기한'], errors='coerce')
        
        merged_df = pd.merge(df_schedule, df_master[['가산관리번호', '명칭', '출원인', '진행현황']], on='가산관리번호', how='left')
        
        return df_master, merged_df
    except Exception as e:
        st.error(f"❌ 데이터 로드 오류: {e}")
        return pd.DataFrame(), pd.DataFrame()

# --- 4. 메인 UI 및 로직 ---
st.sidebar.title("📂 특허 관리 도구")

if st.sidebar.button("🔄 데이터 새로고침"):
    st.cache_data.clear()
    st.rerun()

df_master, df_all = load_data()

if df_master.empty:
    st.warning("데이터 파일({PATENT_FILE})을 확인해주세요.")
    st.stop()

# --- 💾 엑셀 저장 공통 함수 (openpyxl 사용 - 서식 보존) ---
def save_to_excel(new_row_dict, sheet_name):
    try:
        wb = openpyxl.load_workbook(PATENT_FILE)
        if sheet_name not in wb.sheetnames:
            return False, f"❌ 시트를 찾을 수 없습니다: {sheet_name}"
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        row_values = []
        for header in headers:
            val = new_row_dict.get(header)
            row_values.append(val)
        ws.append(row_values)
        wb.save(PATENT_FILE)
        return True, "저장되었습니다."
    except PermissionError:
        return False, "❌ 엑셀 파일이 열려있습니다. 파일을 닫고 다시 시도해주세요."
    except Exception as e:
        return False, f"❌ 저장 중 오류 발생: {e}"

# --- ✏️ 엑셀 수정 함수 (특허) ---
def update_patent_in_excel(target_id, updated_data):
    try:
        wb = openpyxl.load_workbook(PATENT_FILE)
        ws = wb['기본정보(Master)']
        headers = [cell.value for cell in ws[1]]
        
        id_col_idx = -1
        for i, h in enumerate(headers):
            if h == '가산관리번호':
                id_col_idx = i
                break
        if id_col_idx == -1: return False, "가산관리번호 컬럼을 찾을 수 없습니다."

        target_row_idx = -1
        for r in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=r, column=id_col_idx + 1).value
            if str(cell_val) == str(target_id):
                target_row_idx = r
                break
        if target_row_idx == -1: return False, "대상 특허를 찾을 수 없습니다."
        
        for col_name, new_val in updated_data.items():
            if col_name in headers:
                col_idx = headers.index(col_name) + 1
                ws.cell(row=target_row_idx, column=col_idx).value = new_val
        wb.save(PATENT_FILE)
        return True, "수정되었습니다."
    except Exception as e:
        return False, f"수정 오류: {e}"

# --- 🗑️ 엑셀 삭제 함수 (특허) ---
def delete_patent_from_excel(target_id):
    try:
        wb = openpyxl.load_workbook(PATENT_FILE)
        ws = wb['기본정보(Master)']
        headers = [cell.value for cell in ws[1]]
        
        id_col_idx = -1
        for i, h in enumerate(headers):
            if h == '가산관리번호':
                id_col_idx = i
                break
        if id_col_idx == -1: return False, "가산관리번호 컬럼을 찾을 수 없습니다."
        
        target_row_idx = -1
        for r in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=r, column=id_col_idx + 1).value
            if str(cell_val) == str(target_id):
                target_row_idx = r
                break
        if target_row_idx == -1: return False, "대상 특허를 찾을 수 없습니다."
        
        ws.delete_rows(target_row_idx)
        wb.save(PATENT_FILE)
        return True, "삭제되었습니다."
    except Exception as e:
        return False, f"삭제 오류: {e}"

# --- 🔢 비용 ID 생성 함수 ---
def generate_cost_id():
    today_str = datetime.now().strftime("%Y%m%d")
    base_id = f"C-{today_str}-"
    try:
        df_sch = pd.read_excel(PATENT_FILE, sheet_name='비용관리(Schedule)', engine='openpyxl')
        max_seq = 0
        if '비용관리번호' in df_sch.columns:
            existing_ids = df_sch['비용관리번호'].dropna().astype(str)
            target_ids = existing_ids[existing_ids.str.startswith(base_id)]
            for cid in target_ids:
                try:
                    seq = int(cid.split('-')[-1])
                    if seq > max_seq: max_seq = seq
                except: pass
        new_seq = max_seq + 1
        return f"{base_id}{new_seq:02d}"
    except:
        return f"{base_id}01"

# --- ✏️ 비용 수정 함수 ---
def update_cost_in_excel(cost_id, updated_data):
    try:
        wb = openpyxl.load_workbook(PATENT_FILE)
        ws = wb['비용관리(Schedule)']
        headers = [cell.value for cell in ws[1]]
        
        id_col_idx = -1
        for i, h in enumerate(headers):
            if h == '비용관리번호':
                id_col_idx = i
                break
        if id_col_idx == -1: return False, "비용관리번호 컬럼을 찾을 수 없습니다."

        target_row_idx = -1
        for r in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=r, column=id_col_idx + 1).value
            if str(cell_val) == str(cost_id):
                target_row_idx = r
                break
        if target_row_idx == -1: return False, "대상 비용 항목을 찾을 수 없습니다."
        
        for col_name, new_val in updated_data.items():
            if col_name in headers:
                col_idx = headers.index(col_name) + 1
                ws.cell(row=target_row_idx, column=col_idx).value = new_val
        wb.save(PATENT_FILE)
        return True, "비용 정보가 수정되었습니다."
    except Exception as e:
        return False, f"비용 수정 오류: {e}"

# --- 🗑️ 비용 삭제 함수 ---
def delete_cost_from_excel(cost_id):
    try:
        wb = openpyxl.load_workbook(PATENT_FILE)
        ws = wb['비용관리(Schedule)']
        headers = [cell.value for cell in ws[1]]
        
        id_col_idx = -1
        for i, h in enumerate(headers):
            if h == '비용관리번호':
                id_col_idx = i
                break
        if id_col_idx == -1: return False, "비용관리번호 컬럼을 찾을 수 없습니다."
        
        target_row_idx = -1
        for r in range(2, ws.max_row + 1):
            cell_val = ws.cell(row=r, column=id_col_idx + 1).value
            if str(cell_val) == str(cost_id):
                target_row_idx = r
                break
        if target_row_idx == -1: return False, "대상 비용 항목을 찾을 수 없습니다."
        
        ws.delete_rows(target_row_idx)
        wb.save(PATENT_FILE)
        return True, "비용 항목이 삭제되었습니다."
    except Exception as e:
        return False, f"비용 삭제 오류: {e}"

# 메뉴 설정 (Key-Value Mapping)
menu_items = {
    "dashboard": "**📊 대시보드**", # Bold 적용
    "monthly": "📅 월별 비용/업무 관리",
    "detail": "🔍 특허별 상세 이력",
    "patent_input": "📝 데이터 입력/수정(특허)",
    "cost_input": "💰 데이터 입력/수정(비용)"
}

view_mode = st.sidebar.radio(
    "조회 모드", 
    list(menu_items.keys()), 
    format_func=lambda x: menu_items[x]
)

# ==========================================
# [모드 0] 대시보드 (Dashboard)
# ==========================================
# [모드 0] 대시보드 (Dashboard)
# ==========================================
if view_mode == "dashboard":
    st.title("📊 통합 대시보드")
    st.info("특허/지식재산권 현황과 비용 추이를 한눈에 확인하세요.")

    if df_master.empty:
        st.warning("등록된 특허 데이터가 없습니다. '데이터 입력' 메뉴에서 먼저 데이터를 등록해주세요.")
    else:
        # --- 1. 상단 핵심 지표 (Metrics) ---
        total_cnt = len(df_master)
        
        # 카테고리별 집계
        cat_counts = df_master['카테고리'].value_counts()
        p_cnt = cat_counts.get('특허', 0)
        d_cnt = cat_counts.get('디자인', 0)
        t_cnt = cat_counts.get('상표', 0)
        
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("전체 건수", f"{total_cnt}건")
        m2.metric("특허", f"{p_cnt}건")
        m3.metric("디자인", f"{d_cnt}건")
        m4.metric("상표", f"{t_cnt}건")
        
        st.divider()

        # --- 2. 그래프 영역 (2단 레이아웃) ---
        c1, c2 = st.columns(2)
        
        with c1:
            st.subheader("📅 연도별 출원/등록")
            # 연도 추출 및 집계 (출원일자, 등록일자)
            df_chart = df_master.copy()
            
            # 출원 데이터
            app_counts = pd.to_datetime(df_chart['출원일자'], errors='coerce').dt.year.value_counts().reset_index()
            app_counts.columns = ['Year', 'Count']
            app_counts['Type'] = '출원'
            
            # 등록 데이터
            reg_counts = pd.to_datetime(df_chart['등록일자'], errors='coerce').dt.year.value_counts().reset_index()
            reg_counts.columns = ['Year', 'Count']
            reg_counts['Type'] = '등록'
            
            # 데이터 병합
            combined_counts = pd.concat([app_counts, reg_counts]).dropna()
            combined_counts['Year'] = combined_counts['Year'].astype(int)
            
            if not combined_counts.empty:
                # Altair Grouped Bar Chart (Using xOffset for cleaner layout)
                # column(facet) 방식은 화면 너비를 과도하게 차지하여 깨질 수 있음 -> xOffset 방식 사용
                chart_year = alt.Chart(combined_counts).mark_bar(cornerRadius=5, size=20).encode(
                    x=alt.X('Year:O', title='연도'), 
                    y=alt.Y('Count:Q', title='건수'),
                    color=alt.Color('Type:N', scale=alt.Scale(domain=['출원', '등록'], range=['#36a2eb', '#4bc0c0']), legend=alt.Legend(title="구분", orient="top")),
                    xOffset='Type:N', # Grouped alignment (Side-by-side)
                    tooltip=['Year', 'Type', 'Count']
                ).properties(height=300)
                
                st.altair_chart(chart_year, width="stretch")
            else:
                st.caption("출원/등록 일자 데이터가 없습니다.")

        with c2:
            st.subheader("📊 상태별 분포 (비중)")
            # Altair Donut Chart
            status_counts = df_master['진행현황'].value_counts().reset_index()
            status_counts.columns = ['진행현황', 'count']
            
            if not status_counts.empty:
                base = alt.Chart(status_counts).encode(
                    theta=alt.Theta("count", stack=True)
                )
                
                # [수정] 상태별 색상 지정 (일관된 스타일 적용)
                status_domain = ["등록", "심사중", "공고", "출원", "거절", "포기"]
                status_range = ["#28a745", "#ffc107", "#17a2b8", "#007bff", "#dc3545", "#6c757d"]
                
                pie = base.mark_arc(innerRadius=50).encode(
                    color=alt.Color("진행현황", 
                        scale=alt.Scale(domain=status_domain, range=status_range),
                        legend=alt.Legend(title="상태")
                    ),
                    order=alt.Order("count", sort="descending"),
                    tooltip=["진행현황", "count"]
                )
                
                # 텍스트 레이블 (선택사항, 깔끔함을 위해 제외하거나 추가 가능)
                st.altair_chart(pie, width="stretch")
            else:
                st.caption("데이터가 없습니다.")
        
        st.divider()
        
        # --- 3. 비용 및 긴급 업무 ---
        c3, c4 = st.columns([2, 1])
        
        with c3:
            st.subheader("💰 연도별/월별 비용")
            if df_all.empty:
                st.caption("비용 데이터가 없습니다.")
            else:
                df_cost_chart = df_all.copy()
                # [수정] Altair에서 시간 축(:T)을 정확히 인식하도록 String 대신 Timestamp(매월 1일)로 변환
                df_cost_chart = df_cost_chart.dropna(subset=['기한(Due Date)']) # NaT 제거
                if not df_cost_chart.empty:
                    df_cost_chart['년월'] = df_cost_chart['기한(Due Date)'].dt.to_period('M').dt.to_timestamp()
                    df_cost_chart['Year'] = df_cost_chart['기한(Due Date)'].dt.year
                    
                    # 그래프 위치 확보 (필터를 아래에 두기 위함)
                    chart_placeholder = st.empty()
                    
                    # 연도 선택 필터 구성 (Range)
                    current_year = pd.Timestamp.now().year
                    data_years = sorted(df_cost_chart['Year'].unique())
                    
                    # [수정] 사용자가 요청한 고정 범위 (2017 ~ 현재+5년)와 실제 데이터 연도 병합
                    requested_range = range(2017, current_year + 6)
                    available_years = sorted(list(set(data_years) | set(requested_range)))
                    
                    if available_years:
                        # 기본값 설정: (현재년도-2) ~ (현재년도+2)
                        default_start = current_year - 2
                        default_end = current_year + 2
                        
                        # 실제 데이터 범위 내에서 근사값 찾기
                        # available_years는 이제 넓은 범위를 가지므로 min/max 대신 값 존재 여부만 체크하면 됨
                        start_year_opt = default_start if default_start in available_years else min(available_years)
                        end_year_opt = default_end if default_end in available_years else max(available_years)
                        
                        # Selectbox 2개 배치
                        f_col1, f_col2 = st.columns(2)
                        with f_col1:
                            start_year = st.selectbox("시작 연도", available_years, 
                                                    index=available_years.index(start_year_opt) if start_year_opt in available_years else 0,
                                                    key="dashboard_start_year")
                        with f_col2:
                            # 끝 연도가 시작 연도보다 작으면 자동으로 맞춰주거나 별도 처리 (여기선 사용자 선택 맡김)
                            end_year = st.selectbox("종료 연도", available_years, 
                                                  index=available_years.index(end_year_opt) if end_year_opt in available_years else len(available_years)-1,
                                                  key="dashboard_end_year")
                            
                        # 필터링 로직 (Start ~ End Inclusive)
                        # 사용자가 start > end로 선택했을 경우 자동 스왑 처리
                        r_start = min(start_year, end_year)
                        r_end = max(start_year, end_year)
                        
                        df_filtered = df_cost_chart[
                            (df_cost_chart['Year'] >= r_start) & 
                            (df_cost_chart['Year'] <= r_end)
                        ]
                        
                        # 월별 비용 합계 계산
                        monthly_cost = df_filtered.groupby('년월')['금액'].sum().reset_index()
                        
                        if not monthly_cost.empty:
                            chart_cost = alt.Chart(monthly_cost).mark_bar(cornerRadius=5).encode(
                                x=alt.X('년월:T', title='기간', axis=alt.Axis(format='%Y-%m')),
                                y=alt.Y('금액:Q', title='비용'),
                                color=alt.value('#4bc0c0'),
                                tooltip=[alt.Tooltip('년월:T', format='%Y-%m'), alt.Tooltip('금액', format=',.0f')]
                            ).properties(height=300)
                            
                            chart_placeholder.altair_chart(chart_cost, width="stretch")
                        else:
                            chart_placeholder.info(f"{r_start}년 ~ {r_end}년 비용 내역이 없습니다.")
                    else:
                         chart_placeholder.caption("표시할 연도 데이터가 없습니다.")
                else:
                    st.caption("기간 정보가 있는 비용 데이터가 없습니다.")
                
        with c4:
            st.subheader("🚨 긴급 회신/납부 항목 (Top 5)")
            today = pd.Timestamp.now()
            
            if df_all.empty:
                st.caption("예정된 업무가 없습니다.")
            else:
                # 기한이 오늘 이후인 것 중 가장 가까운 것
                urgent_tasks = df_all[
                    (df_all['기한(Due Date)'] >= today) 
                ].sort_values('기한(Due Date)').head(5)
                
                if urgent_tasks.empty:
                    st.success("예정된 긴급 업무가 없습니다! 🎉")
                else:
                    for _, row in urgent_tasks.iterrows():
                        d_day = (row['기한(Due Date)'] - today).days
                        d_str = "오늘" if d_day == 0 else f"D-{d_day}" if d_day > 0 else f"D+{abs(d_day)}"
                        
                        with st.container(border=True):
                            st.markdown(f"**{row['구분']}** ({d_str})")
                            st.caption(f"{row['명칭']} | {row['기한(Due Date)'].strftime('%Y-%m-%d')}")
                            if pd.notna(row['금액']):
                                st.caption(f"💰 {row['금액']:,} {row['통화']}")

# ==========================================
# [모드 1] 월별 비용/업무 관리
# ==========================================
elif view_mode == "monthly":
    st.title("📅 월별 납부 및 업무 현황")

    col1, col2 = st.sidebar.columns(2)
    today = datetime.now()
    
    years = sorted(df_all['기한(Due Date)'].dt.year.dropna().unique().astype(int))
    if not years: years = [today.year]
    selected_year = col1.selectbox("연도", years, index=years.index(today.year) if today.year in years else 0)
    
    month_options = {0: "전체(All)"}
    for m in range(1, 13):
        month_options[m] = f"{m}월"
    
    default_month_idx = today.month
    selected_month_key = col2.selectbox("월", list(month_options.keys()), format_func=lambda x: month_options[x], index=default_month_idx)

    year_mask = (df_all['기한(Due Date)'].dt.year == selected_year)
    
    if selected_month_key != 0:
        target_date_start = datetime(selected_year, selected_month_key, 1)
        if selected_month_key == 12:
            next_month_start = datetime(selected_year + 1, 1, 1)
        else:
            next_month_start = datetime(selected_year, selected_month_key + 1, 1)
        
        month_mask = (df_all['기한(Due Date)'].dt.month == selected_month_key)
        
        overdue_mask = (df_all['납부상태'] == '미납') & (df_all['기한(Due Date)'] < target_date_start)
        current_month_task_mask = (year_mask & month_mask & (df_all['납부상태'] == '미납'))
        
        todo_df = df_all[overdue_mask | current_month_task_mask].copy()
        future_df = df_all[(df_all['납부상태'] == '미납') & (df_all['기한(Due Date)'] >= next_month_start)].copy()
        done_df = df_all[year_mask & month_mask & (df_all['납부상태'] == '납부완료')].copy()
        display_title = f"{selected_year}년 {selected_month_key}월 현황"

    else:
        todo_df = df_all[year_mask & (df_all['납부상태'] == '미납')].copy()
        future_df = df_all[(df_all['기한(Due Date)'].dt.year > selected_year) & (df_all['납부상태'] == '미납')].copy()
        done_df = df_all[year_mask & (df_all['납부상태'] == '납부완료')].copy()
        display_title = f"{selected_year}년 전체 현황"

    st.markdown(f"### 📊 {display_title}")
    
    krw_total = todo_df[todo_df['통화'] == 'KRW']['금액'].sum()
    usd_total = todo_df[todo_df['통화'] == 'USD']['금액'].sum()
    
    # [추가] 납부 완료 금액 계산
    done_krw = done_df[done_df['통화'] == 'KRW']['금액'].sum()
    done_usd = done_df[done_df['통화'] == 'USD']['금액'].sum()
    
    # 레이아웃 변경 (4열 -> 6열 또는 적절히 배치)
    # 예정 금액 | 완료 금액 | 건수 현황
    m1, m2, m3, m4, m5, m6 = st.columns(6)
    m1.metric("납부 예정 (KRW)", f"{krw_total:,.0f} 원")
    m2.metric("납부 예정 (USD)", f"$ {usd_total:,.2f}")
    m3.metric("납부 완료 (KRW)", f"{done_krw:,.0f} 원") # New
    m4.metric("납부 완료 (USD)", f"$ {done_usd:,.2f}") # New
    m5.metric("🚨 처리 대상", f"{len(todo_df)} 건")
    m6.metric("✅ 처리 완료", f"{len(done_df)} 건")

    st.divider()

    tab1, tab2, tab3 = st.tabs(["🚨 처리 대상 (Todo)", "🔮 향후 예정 (Future)", "✅ 납부 완료 (Done)"])
    cols = ['기한(Due Date)', '회신요청기한', '가산관리번호', '명칭', '구분', '금액', '통화', '비고'] if '회신요청기한' in df_all.columns else ['기한(Due Date)', '가산관리번호', '명칭', '구분', '금액', '통화', '비고']

    with tab1:
        st.caption("※ 과거 연체 건과 선택한 기간의 미납 내역입니다.")
        if todo_df.empty:
            st.success("처리해야 할 업무가 없습니다! 🎉")
        else:
            todo_df = todo_df.sort_values('기한(Due Date)')
            def highlight_overdue(row):
                if pd.isnull(row['기한(Due Date)']): return [''] * len(row)
                if row['기한(Due Date)'] < datetime.now():
                    return ['background-color: #ffe6e6; color: #b30000; font-weight: bold'] * len(row)
                return [''] * len(row)

            # [수정] use_container_width -> width="stretch" 로 변경
            st.dataframe(
                todo_df[cols].style.apply(highlight_overdue, axis=1)
                .format({'기한(Due Date)': '{:%Y-%m-%d}', '회신요청기한': '{:%Y-%m-%d}', '금액': '{:,.0f}'}, na_rep="-"),
                width="stretch", 
                hide_index=True
            )

    with tab2:
        if future_df.empty: st.info("예정된 내역이 없습니다.")
        else:
            # [수정] use_container_width -> width="stretch"
            st.dataframe(
                future_df[cols].sort_values('기한(Due Date)').style.format({'기한(Due Date)': '{:%Y-%m-%d}', '회신요청기한': '{:%Y-%m-%d}', '금액': '{:,.0f}'}, na_rep="-"), 
                width="stretch", 
                hide_index=True
            )

    with tab3:
        if done_df.empty: st.info("완료된 내역이 없습니다.")
        else:
            done_cols = ['실제납부일', '가산관리번호', '명칭', '구분', '금액', '통화', '비고']
            
            # [핵심 수정] .style 제거하고 column_config 사용 (에러 해결)
            st.dataframe(
                done_df[done_cols].sort_values('실제납부일', ascending=False),
                column_config={
                    "실제납부일": st.column_config.DateColumn("실제납부일", format="YYYY-MM-DD"),
                    "금액": st.column_config.NumberColumn("금액", format="%d"),
                },
                width="stretch", # 경고 해결
                hide_index=True
            )

# ==========================================
# [모드 2] 특허별 상세 이력
# ==========================================
# [모드 2] 특허별 상세 이력
# ==========================================
elif view_mode == "detail":
    st.title("🔍 특허 상세 이력 조회")
    
    search_options = df_master.apply(lambda x: f"[{x['가산관리번호']}] {x['명칭']}", axis=1)
    selected_option = st.selectbox("특허 검색", search_options)
    
    selected_id = selected_option.split(']')[0].replace('[', '')
    
    target_master = df_master[df_master['가산관리번호'] == selected_id].iloc[0]
    target_schedule = df_all[df_all['가산관리번호'] == selected_id].copy()
    target_schedule['비용단계'] = target_schedule['구분'].apply(classify_cost_stage)

    with st.expander("ℹ️ 기본 정보 펼치기", expanded=True):
        c1, c2 = st.columns([3, 1])
        with c1:
            st.markdown(f"### {target_master['명칭']}")
            st.caption("관리번호 (Copy 가능)")
            st.code(target_master['가산관리번호'], language=None)
            st.caption(f"출원인: {target_master['출원인']}")
        with c2:
            st.markdown(f"**현재상태: :blue[{target_master['진행현황']}]**")
        
        st.divider()
        
        info_c1, info_c2, info_c3 = st.columns(3)
        with info_c1:
            st.markdown("**출원번호**")
            val = str(target_master['출원번호']) if pd.notna(target_master['출원번호']) else "-"
            st.code(val, language=None)
        with info_c2:
            st.markdown("**출원일자**")
            val = format_date_clean(target_master['출원일자'])
            st.code(val, language=None)
        with info_c3:
            st.markdown("**등록번호**")
            val = str(target_master['등록번호']) if pd.notna(target_master['등록번호']) else "-"
            st.code(val, language=None)
        
        if pd.notna(target_master['설명']):
            st.caption("설명")
            st.info(target_master['설명'])

    st.subheader("💰 비용 단계별 요약")
    
    summary = target_schedule.groupby(['비용단계', '통화'])['금액'].sum().reset_index()
    
    cost_cols = st.columns(3)
    stages = ['출원 단계', '등록 단계', '유지관리(연차료)']
    
    for i, stage in enumerate(stages):
        with cost_cols[i]:
            with st.container(border=True):
                st.markdown(f"##### {stage}")
                stage_sum = summary[summary['비용단계'] == stage]
                
                if stage_sum.empty:
                    st.metric(label="합계", value="0")
                else:
                    for _, row in stage_sum.iterrows():
                        currency_label = "KRW (원)" if row['통화'] == 'KRW' else "USD ($)"
                        st.metric(
                            label=currency_label, 
                            value=f"{row['금액']:,.0f}" if row['통화'] == 'KRW' else f"{row['금액']:,.2f}"
                        )

    st.divider()

    st.subheader("📋 전체 상세 내역")
    display_cols = ['비용단계', '기한(Due Date)', '회신요청기한', '구분', '금액', '통화', '납부상태', '실제납부일', '비고']
    if '회신요청기한' not in target_schedule.columns:
        display_cols.remove('회신요청기한')
    
    def color_status(val):
        return 'background-color: #d4edda; color: #155724' if val == '납부완료' else 'background-color: #fff3cd; color: #856404'

    # [수정] use_container_width -> width="stretch"
    st.dataframe(
        target_schedule[display_cols].sort_values('기한(Due Date)').style
        .map(color_status, subset=['납부상태'])
        .format({'기한(Due Date)': '{:%Y-%m-%d}', '회신요청기한': '{:%Y-%m-%d}', '금액': '{:,.0f}'}, na_rep="-"),
        width="stretch",
        hide_index=True
    )

# ==========================================
# [모드 3] 데이터 입력/수정(특허)
# ==========================================
# [모드 3] 데이터 입력/수정(특허)
# ==========================================
elif view_mode == "patent_input":
    st.title("📝 데이터 입력/수정(특허)")
    st.info("특허 및 지식재산권 기본 정보를 등록하거나 수정합니다.")

    tab_input1, tab_input3, tab_input2 = st.tabs(["🆕 신규 특허 등록", "🛠️ 특허 수정/삭제", "📋 전체 특허 목록"])

    # --- 1. 신규 특허 등록 폼 ---
    with tab_input1:
        st.subheader("새로운 특허/지식재산권 등록")
        
        # [수정] 엑셀 컬럼 기반 동적 폼 생성
        current_columns = list(df_master.columns)
        input_data = {}

        with st.form("master_form"):
            # 2열 레이아웃 적용을 위한 준비
            cols = st.columns(2)
            
            for i, col_name in enumerate(current_columns):
                # 컬럼마다 위치할 col 지정 (0 또는 1)
                with cols[i % 2]:
                    # 1. 날짜 필드 감지
                    if '일자' in col_name or 'date' in col_name.lower():
                        input_data[col_name] = st.date_input(col_name, value=None)
                    
                    # 2. 특정 필드 커스텀 (진행현황)
                    elif col_name == '진행현황':
                        input_data[col_name] = st.selectbox(col_name, ["-", "출원", "심사중", "공고", "등록", "거절", "포기"])
                        
                    # 3. 설명/메모 필드 (Text Area)
                    elif col_name == '설명':
                        input_data[col_name] = st.text_area(col_name)
                        
                    # 4. 그 외 텍스트 필드
                    else:
                        # 필수값 표시
                        label = col_name
                        if col_name == '가산관리번호':
                            label += " (Key/필수)"
                            input_data[col_name] = st.text_input(label, placeholder="예: P-24-001")
                        elif col_name == '명칭':
                            label += " (필수)"
                            input_data[col_name] = st.text_input(label)
                        else:
                            input_data[col_name] = st.text_input(label)
                            
            submitted = st.form_submit_button("신규 특허 저장")
            
            if submitted:
                # 필수값 검증
                manage_no = input_data.get('가산관리번호')
                name = input_data.get('명칭')
                
                if not manage_no or not name:
                    st.error("가산관리번호와 명칭은 필수 입력값입니다.")
                else:
                    # 중복 체크
                    if manage_no in df_master['가산관리번호'].values:
                        st.error("이미 존재하는 관리번호입니다.")
                    else:
                        # 저장할 데이터 (input_data 그대로 사용)
                        success, pid_msg = save_to_excel(input_data, '기본정보(Master)')
                        if success:
                            st.success(pid_msg)
                            time.sleep(1)
                            st.cache_data.clear()
                            st.rerun()
                        else:
                            st.error(pid_msg)

        st.divider()
        st.subheader("📋 현재 등록된 특허 목록")
        
        df_display = df_master.copy()
        if '카테고리' in df_display.columns:
            df_display['카테고리'] = df_display['카테고리'].apply(lambda x: [x] if pd.notna(x) else [])

        def style_patent_status(val):
            if val == '등록':
                return 'background-color: #d4edda; color: #155724; font-weight: bold' # Green
            elif val == '심사중':
                return 'background-color: #fff3cd; color: #856404' # Yellow
            elif val == '출원':
                return 'background-color: #cce5ff; color: #004085' # Blue
            elif val == '공고':
                return 'background-color: #d1ecf1; color: #0c5460' # Cyan
            elif val == '거절':
                return 'background-color: #f8d7da; color: #721c24' # Red
            elif val == '포기':
                return 'background-color: #e2e3e5; color: #383d41' # Gray
            return ''

        st.dataframe(
            df_display.sort_values('가산관리번호').style.map(style_patent_status, subset=['진행현황']),
            width="stretch",
            hide_index=True
        )

    # --- 2. 수정/삭제 폼 (특허용) ---
    with tab_input3:
        st.subheader("등록된 특허 수정/삭제")
        
        # 대상 선택
        edit_options = df_master.apply(lambda x: f"[{x['가산관리번호']}] {x['명칭']}", axis=1)
        selected_edit_item = st.selectbox("수정/삭제할 특허 선택", edit_options, key="edit_select_patent")
        
        if selected_edit_item:
            target_edit_id = selected_edit_item.split(']')[0].replace('[', '')
            
            # 현재 선택된 특허의 원본 데이터 가져오기
            current_row = df_master[df_master['가산관리번호'] == target_edit_id].iloc[0]
            
            st.divider()
            
            # --- 수정 모드 ---
            st.markdown("##### ✏️ 내용 수정")
            
            edit_data = {}
            edit_cols = list(df_master.columns)
            
            with st.form("edit_form_patent"):
                e_cols = st.columns(2)
                for i, col_name in enumerate(edit_cols):
                    with e_cols[i % 2]:
                        # 기존 값 가져오기
                        val = current_row[col_name]
                        if pd.isna(val): val = ""
                        
                        # 가산관리번호는 수정 불가 (Primary Key)
                        if col_name == '가산관리번호':
                            st.text_input(col_name, value=val, disabled=True)
                            edit_data[col_name] = val # 값은 유지
                        
                        elif '일자' in col_name or 'date' in col_name.lower():
                            d_val = None
                            if str(val).strip() != "":
                                try: d_val = pd.to_datetime(val).date()
                                except: pass
                            edit_data[col_name] = st.date_input(f"{col_name} (수정)", value=d_val)
                            
                        elif col_name == '진행현황':
                            status_opts = ["-", "출원", "심사중", "공고", "등록", "거절", "포기"]
                            try: idx = status_opts.index(val)
                            except: idx = 0
                            edit_data[col_name] = st.selectbox(f"{col_name} (수정)", status_opts, index=idx)
                            
                        elif col_name == '설명':
                            edit_data[col_name] = st.text_area(f"{col_name} (수정)", value=str(val))
                            
                        else:
                            edit_data[col_name] = st.text_input(f"{col_name} (수정)", value=str(val))
            
                update_submitted = st.form_submit_button("수정 내용 저장")
                
                if update_submitted:
                    success, msg = update_patent_in_excel(target_edit_id, edit_data)
                    if success:
                        st.success(msg)
                        time.sleep(1)
                        st.cache_data.clear()
                        st.rerun()
                    else:
                        st.error(msg)
            
            st.divider()
            
            # --- 삭제 모드 ---
            st.markdown("##### 🗑️ 특허 삭제")
            st.warning("주의: '기본정보(Master)' 시트에서만 삭제됩니다. '비용관리' 시트의 내역은 자동으로 삭제되지 않습니다.")
            
            confirm_delete = st.checkbox("정말로 이 특허를 삭제하시겠습니까?")
            
            if st.button("선택한 특허 삭제", type="primary", disabled=not confirm_delete):
                success, msg = delete_patent_from_excel(target_edit_id)
                if success:
                    st.success(msg)
                    time.sleep(1)
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error(msg)

# ==========================================
    # --- 3. 전체 특허 목록 탭 ---
    with tab_input2:
        st.subheader("📋 전체 특허 목록")
        st.info("현재 등록된 모든 특허/지식재산권 목록입니다.")
        
        df_display_all = df_master.copy()
        if not df_display_all.empty:
            st.dataframe(
                df_display_all.sort_values('가산관리번호').style.map(style_patent_status, subset=['진행현황']),
                width="stretch",
                hide_index=True
            )
        else:
            st.info("등록된 특허가 없습니다.")

# ==========================================
# [모드 4] 데이터 입력/수정(비용)
# ==========================================
# [모드 4] 데이터 입력/수정(비용)
# ==========================================
elif view_mode == "cost_input":
    st.title("💰 데이터 입력/수정(비용)")
    st.info("비용 및 일정 정보를 등록하거나 수정합니다.")

    tab_cost1, tab_cost2, tab_cost3 = st.tabs(["💰 비용/일정 추가", "🛠️ 비용 수정/삭제", "📋 전체 비용 목록"])

    # --- 1. 비용/일정 추가 폼 ---
    with tab_cost1:
        st.subheader("기존 특허에 비용/일정 추가")
        
        if df_master.empty:
            st.error("등록된 특허가 없습니다. 먼저 특허를 등록해주세요.")
        else:
            mp_options = df_master.apply(lambda x: f"[{x['가산관리번호']}] {x['명칭']}", axis=1)
            selected_mp = st.selectbox("대상 특허 선택", mp_options)
            target_id = selected_mp.split(']')[0].replace('[', '')
            
            # 자동 생성된 ID (미리보기)
            next_id = generate_cost_id()
            
            with st.form("schedule_form"):
                st.info(f"💡 비용관리번호가 자동 생성됩니다: {next_id}")
                
                s_col1, s_col2 = st.columns(2)
                cost_name = s_col1.text_input("구분 (비용명)", placeholder="예: 4년차 연차료")
                due_date = s_col2.date_input("납부기한 (Due Date)")
                
                # 회신요청기한 추가
                reply_date = st.date_input("회신요청기한", value=None)
                
                s_col3, s_col4, s_col5 = st.columns(3)
                amount = s_col3.number_input("금액", min_value=0, step=1000)
                currency = s_col4.selectbox("통화", ["KRW", "USD"])
                pay_status = s_col5.selectbox("납부상태", ["미납", "납부완료"])
                
                pay_date = st.date_input("실제납부일 (납부완료 시)", value=None)
                note = st.text_input("비고")
                
                sch_submitted = st.form_submit_button("비용 일정 저장")
                
                if sch_submitted:
                    if not cost_name:
                        st.error("비용 구분을 입력해주세요.")
                    else:
                        new_sch_data = {
                            "비용관리번호": next_id, # 자동생성 ID 추가
                            "가산관리번호": target_id,
                            "구분": cost_name,
                            "기한(Due Date)": due_date,
                            "회신요청기한": reply_date, # 추가
                            "금액": amount,
                            "통화": currency,
                            "납부상태": pay_status,
                            "실제납부일": pay_date if pay_status == '납부완료' else None,
                            "비고": note
                        }
                        success, msg = save_to_excel(new_sch_data, '비용관리(Schedule)')
                        if success:
                            st.success(msg)
                            time.sleep(1)
                            st.cache_data.clear()
                            st.rerun()
                        else:
                            st.error(msg)
            
            st.divider()
            st.markdown(f"##### 📊 현재 등록된 비용/일정 내역 ({target_id})")
            
            if not df_all.empty:
                related_costs = df_all[df_all['가산관리번호'] == target_id].copy()
                
                if related_costs.empty:
                    st.info("등록된 비용 내역이 없습니다.")
                else:
                    # 비용관리번호 컬럼이 없어도 에러 안 나게 처리
                    cols_to_show = ['비용관리번호', '구분', '기한(Due Date)', '회신요청기한', '금액', '통화', '납부상태', '실제납부일', '비고']
                    # 실제 존재하는 컬럼만 표시
                    final_cols = [c for c in cols_to_show if c in related_costs.columns]

                    def color_status(val):
                        return 'background-color: #d4edda; color: #155724' if val == '납부완료' else 'background-color: #fff3cd; color: #856404'
                    
                    st.dataframe(
                        related_costs[final_cols].sort_values('기한(Due Date)')
                        .style.map(color_status, subset=['납부상태']),
                        width="stretch",
                        hide_index=True,
                        column_config={
                            "기한(Due Date)": st.column_config.DateColumn("기한", format="YYYY-MM-DD"),
                            "회신요청기한": st.column_config.DateColumn("회신요청기한", format="YYYY-MM-DD"),
                            "실제납부일": st.column_config.DateColumn("실제납부일", format="YYYY-MM-DD"),
                            "금액": st.column_config.NumberColumn("금액", format="%d")
                        }
                    )
            else:
                 st.info("로드된 데이터가 없습니다.")

    # --- 2. 비용 수정/삭제 폼 ---
    with tab_cost2:
        st.subheader("등록된 비용/일정 수정/삭제")
        
        # 1. 특허 선택 -> 해당 특허의 비용 목록 조회
        mp_options_edit = df_master.apply(lambda x: f"[{x['가산관리번호']}] {x['명칭']}", axis=1)
        selected_mp_edit = st.selectbox("대상 특허 선택", mp_options_edit, key="cost_edit_mp_select")
        target_mp_edit_id = selected_mp_edit.split(']')[0].replace('[', '')
        
        target_costs = df_all[df_all['가산관리번호'] == target_mp_edit_id].copy()
        
        if target_costs.empty:
            st.info("이 특허에 등록된 비용 내역이 없습니다.")
        else:
            # 2. 비용 항목 선택 (비용관리번호 필수)
            if '비용관리번호' not in target_costs.columns:
                st.error("엑셀 파일에 '비용관리번호' 컬럼이 없습니다. 먼저 컬럼을 추가해주세요.")
            else:
                cost_options = target_costs.apply(
                    lambda x: f"[{x['비용관리번호']}] {x['구분']} ({x['금액']} {x['통화']})", axis=1
                )
                selected_cost_item = st.selectbox("수정할 비용 항목 선택", cost_options, key="cost_item_select_box")
                
                if selected_cost_item:
                    target_cost_id = selected_cost_item.split(']')[0].replace('[', '')
                    
                    # [Fix] 안전한 필터링 (문자열 변환 및 존재 여부 확인)
                    filtered_rows = target_costs[target_costs['비용관리번호'].astype(str) == str(target_cost_id)]
                    
                    if filtered_rows.empty:
                        st.error("선택한 비용 정보를 찾을 수 없습니다.")
                        st.stop()
                    
                    current_cost_row = filtered_rows.iloc[0]
                    
                    st.divider()
                    st.markdown("##### ✏️ 비용 정보 수정")
                    
                    with st.form("cost_edit_form"):
                        # 비용관리번호 (수정 불가)
                        st.text_input("비용관리번호 (ID)", value=target_cost_id, disabled=True)
                        
                        ce_col1, ce_col2 = st.columns(2)
                        # 구분
                        new_cost_name = ce_col1.text_input("구분", value=current_cost_row['구분'])
                        # 기한
                        d_val_due = None
                        try: d_val_due = pd.to_datetime(current_cost_row['기한(Due Date)']).date()
                        except: pass
                        try: d_val_due = pd.to_datetime(current_cost_row['기한(Due Date)']).date()
                        except: pass
                        new_due_date = ce_col2.date_input("납부기한", value=d_val_due)
                        
                        # 회신요청기한 수정
                        d_val_reply = None
                        if '회신요청기한' in current_cost_row.index and pd.notna(current_cost_row['회신요청기한']):
                            try: d_val_reply = pd.to_datetime(current_cost_row['회신요청기한']).date()
                            except: pass
                        new_reply_date = st.date_input("회신요청기한", value=d_val_reply)
                        
                        ce_col3, ce_col4, ce_col5 = st.columns(3)
                        # 금액
                        new_amount = ce_col3.number_input("금액", value=int(current_cost_row['금액']), step=1000)
                        # 통화 (현재 값 인덱스 찾기)
                        curr_opts = ["KRW", "USD"]
                        try: c_idx = curr_opts.index(current_cost_row['통화'])
                        except: c_idx = 0
                        new_currency = ce_col4.selectbox("통화", curr_opts, index=c_idx)
                        # 상태
                        stat_opts = ["미납", "납부완료"]
                        try: s_idx = stat_opts.index(current_cost_row['납부상태'])
                        except: s_idx = 0
                        new_status = ce_col5.selectbox("납부상태", stat_opts, index=s_idx)
                        
                        # 실제납부일
                        d_val_pay = None
                        if pd.notna(current_cost_row['실제납부일']):
                            try: d_val_pay = pd.to_datetime(current_cost_row['실제납부일']).date()
                            except: pass
                        new_pay_date = st.date_input("실제납부일", value=d_val_pay)
                        
                        new_note = st.text_input("비고", value=str(current_cost_row['비고']) if pd.notna(current_cost_row['비고']) else "")
                        
                        cost_update_submitted = st.form_submit_button("비용 정보 수정")
                        
                        if cost_update_submitted:
                            updated_cost_data = {
                                "구분": new_cost_name,
                                "기한(Due Date)": new_due_date,
                                "회신요청기한": new_reply_date,
                                "금액": new_amount,
                                "통화": new_currency,
                                "납부상태": new_status,
                                "실제납부일": new_pay_date if new_status == '납부완료' else None,
                                "비고": new_note
                            }
                            success, msg = update_cost_in_excel(target_cost_id, updated_cost_data)
                            if success:
                                st.success(msg)
                                time.sleep(1)
                                st.cache_data.clear()
                                st.rerun()
                            else:
                                st.error(msg)
    
                    st.divider()
                    st.markdown("##### 🗑️ 비용 항목 삭제")
                    confirm_cost_delete = st.checkbox("정말로 이 비용 항목을 삭제하시겠습니까?", key="del_cost_check")
                    
                    if st.button("선택한 비용 삭제", type="primary", disabled=not confirm_cost_delete):
                        success, msg = delete_cost_from_excel(target_cost_id)
                        if success:
                            st.success(msg)
                            time.sleep(1)
                            st.cache_data.clear()
                            st.rerun()
                        else:
                            st.error(msg)

    # --- 3. 전체 비용 목록 탭 ---
    with tab_cost3:
        st.subheader("📋 전체 비용/일정 목록")
        st.info("등록된 모든 비용 및 일정 내역입니다.")
        
        if df_all.empty:
            st.info("등록된 비용 내역이 없습니다.")
        else:
            # 표시할 컬럼 정의
            cols_all_cost = ['기한(Due Date)', '회신요청기한', '가산관리번호', '명칭', '구분', '금액', '통화', '납부상태', '실제납부일', '비용관리번호']
            # 없는 컬럼 제외 (안전장치)
            final_cols_all = [c for c in cols_all_cost if c in df_all.columns]
            
            st.dataframe(
                df_all[final_cols_all].sort_values('기한(Due Date)', ascending=False),
                width="stretch",
                hide_index=True,
                column_config={
                    "기한(Due Date)": st.column_config.DateColumn("기한", format="YYYY-MM-DD"),
                    "회신요청기한": st.column_config.DateColumn("회신요청기한", format="YYYY-MM-DD"),
                    "실제납부일": st.column_config.DateColumn("실제납부일", format="YYYY-MM-DD"),
                    "금액": st.column_config.NumberColumn("금액", format="%d")
                }
            )
